"""Exports: the formula guard, the columns, and a workbook a colleague can use.

Two things are being defended here. The guard, because a tender title is text a
stranger wrote and a spreadsheet runs a cell that starts with the wrong
character. And the cell types, because an export where every cell is text gives
nonsense when somebody sorts or sums it, and they will blame the tool.
"""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import load_workbook

from watchdog.core.enums import (
    Band,
    ContractNature,
    DeadlineType,
    Domain,
    NoticeStage,
    RuleSignal,
    RulesRoute,
    RuleStrength,
    SourcePlatform,
)
from watchdog.core.models import (
    RegisterRow,
    RuleMatch,
    RulesResult,
    ScreeningResult,
    Tender,
    TextBlock,
)
from watchdog.services import export


def _row(title: str, *, native: str | None = None) -> RegisterRow:
    tender = Tender(
        source=SourcePlatform.TED,
        source_id="900-2026",
        title=title,
        title_native=native,
        title_native_language="deu",
        buyer_name="Stadtwerke",
        buyer_country="DEU",
        place_of_performance_country=["DEU"],
        published_date=date(2026, 9, 1),
        deadline_date=date(2026, 9, 30),
        deadline_type=DeadlineType.PARTICIPATION_REQUEST,
        notice_stage=NoticeStage.CONTRACT_NOTICE,
        contract_nature=ContractNature.SERVICES,
        contract_natures=[ContractNature.SERVICES],
        # A real CPV code with a leading zero: electricity, the one that proves
        # the point about codes that look like numbers.
        cpv_main="09310000",
        cpv_all=["09310000"],
        estimated_value=Decimal("1250000.50"),
        currency="EUR",
        screening_blocks=[TextBlock(field="title-proc", language="deu", text=title)],
    )
    screening = ScreeningResult(
        tender_id=tender.id,
        score=2,
        band=Band.REVIEW,
        rules_only_score=2,
        confidence=0.55,
        reason_codes=["RULES_ONLY"],
        rules=RulesResult(
            tender_id=tender.id,
            matches=[
                RuleMatch(
                    rule_id="grid_transmission",
                    signal=RuleSignal.DOMAIN,
                    strength=RuleStrength.HIGH,
                    alias_matched="umspannwerk",
                    field="title-proc",
                    evidence="Sanierung von Umspannwerken",
                )
            ],
            domains_hit=[Domain.GRID_TRANSMISSION],
            route=RulesRoute.ASSESS,
            rules_version="3",
        ),
        explanation="Rules-only grade 2.",
        screened_content_hash=tender.content_hash,
        provider="disabled",
        rules_version="3",
        policy_version="1",
        profile_version="1",
    )
    return RegisterRow(tender=tender, screening=screening, days_left=17)


@pytest.mark.parametrize("dangerous", ["=1+1", "+SUM(A1)", "-2+3", "@SUM(A1)", "\tcmd", "\rcmd"])
def test_a_cell_a_spreadsheet_would_run_is_prefixed(dangerous: str) -> None:
    assert export.guard(dangerous).startswith("'")
    assert export.guard(dangerous)[1:] == dangerous


@pytest.mark.parametrize("harmless", ["Germany - Engineering", "2026-09-30", "", "Hydrogen"])
def test_an_ordinary_cell_is_left_exactly_as_it_was(harmless: str) -> None:
    assert export.guard(harmless) == harmless


def test_the_guard_reaches_the_csv(fake_title: str = '=HYPERLINK("http://x","click")') -> None:
    body = b"".join(export.to_csv([_row(fake_title)])).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(body)))

    title = rows[1][export.COLUMNS.index("title")]
    assert title.startswith("'=HYPERLINK")


def test_the_guard_reaches_the_workbook() -> None:
    sheet = _sheet([_row("=cmd|' /c calc'!A1")])
    cell = _cell(sheet, "title")

    assert cell.data_type == "s"
    assert str(cell.value).startswith("'=cmd")


def test_source_text_is_exported_exactly_as_it_was_stored() -> None:
    native = "Sanierung von Umspannwerken \u2013 Los 2, Bestandsanlage"
    body = b"".join(export.to_csv([_row("Germany - Works", native=native)])).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(body)))

    assert rows[1][export.COLUMNS.index("title_native")] == native


def test_the_export_says_which_kind_of_number_the_score_is() -> None:
    body = b"".join(export.to_csv([_row("Germany - Works")])).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(body)))

    assert rows[1][export.COLUMNS.index("score_kind")] == "rules_only"


def test_an_unknown_value_exports_as_empty_never_as_none_or_zero() -> None:
    row = _row("Germany - Works")
    row.tender.buyer_name = None
    row.days_left = None

    body = b"".join(export.to_csv([row])).decode("utf-8-sig")
    values = list(csv.reader(io.StringIO(body)))[1]

    assert values[export.COLUMNS.index("buyer_name")] == ""
    assert values[export.COLUMNS.index("days_left")] == ""


def test_an_unknown_value_is_an_empty_cell_in_the_workbook() -> None:
    row = _row("Germany - Works")
    row.tender.buyer_name = None
    row.days_left = None
    row.tender.deadline_date = None

    sheet = _sheet([row])

    assert _cell(sheet, "buyer_name").value is None
    assert _cell(sheet, "days_left").value is None
    assert _cell(sheet, "deadline_date").value is None


# ------------------------------------------------------- types in the workbook


def test_a_date_is_written_as_a_date_so_sorting_works() -> None:
    sheet = _sheet([_row("Germany - Works")])

    assert _cell(sheet, "published_date").value == datetime(2026, 9, 1)
    assert _cell(sheet, "deadline_date").value == datetime(2026, 9, 30)


def test_a_timestamp_is_written_in_utc_and_the_column_says_so() -> None:
    row = _row("Germany - Works")
    row.screening.created_at = datetime(2026, 9, 12, 14, 30, tzinfo=UTC)

    sheet = _sheet([row])

    assert "screened_at_utc" in export.COLUMNS
    assert _cell(sheet, "screened_at_utc").value == datetime(2026, 9, 12, 14, 30)


def test_a_number_is_written_as_a_number_so_summing_works() -> None:
    sheet = _sheet([_row("Germany - Works")])

    for column in ("score", "evidence_grade", "days_left"):
        cell = _cell(sheet, column)
        assert isinstance(cell.value, int), f"{column} arrived as {type(cell.value)}"

    assert _cell(sheet, "days_left").value == 17
    assert _cell(sheet, "confidence").value == pytest.approx(0.55)
    assert _cell(sheet, "estimated_value").value == Decimal("1250000.50")


@pytest.mark.parametrize(
    ("code_column", "expected"),
    [
        ("cpv_main", "09310000"),
        ("buyer_country", "DEU"),
        ("rules_version", "3"),
    ],
)
def test_a_code_stays_text_so_a_leading_zero_survives(code_column: str, expected: str) -> None:
    """09310000 is an identifier. Read as a number it becomes 9,310,000."""
    sheet = _sheet([_row("Germany - Works")])
    cell = _cell(sheet, code_column)

    assert cell.value == expected
    assert cell.data_type == "s"
    assert cell.number_format == "@"


def test_the_header_row_is_frozen_and_filterable() -> None:
    sheet = _sheet([_row("Germany - Works")])

    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None


def test_the_workbook_is_a_readable_zip_openpyxl_can_reopen() -> None:
    workbook = export.to_xlsx([_row("Germany - Works")])

    with zipfile.ZipFile(io.BytesIO(workbook)) as archive:
        assert archive.testzip() is None

    reopened = load_workbook(io.BytesIO(workbook))
    assert reopened.sheetnames == ["Register"]
    # The heading row is the team's word for each column, not the internal key.
    assert [cell.value for cell in reopened["Register"][1]] == export.headings()


def test_the_heading_row_uses_the_teams_word_and_the_column_keys_stay_stable() -> None:
    """One display vocabulary, and a label change is never a data change.

    The key is what a stored extraction, a test and any downstream script address;
    the heading is what a colleague reads. Renaming a label must not move a column
    or rename a key.
    """
    headings = export.headings()

    assert len(headings) == len(export.COLUMNS)
    assert headings[export.COLUMNS.index("buyer_name")] == "Client"
    assert headings[export.COLUMNS.index("buyer_country")] == "Client Country"
    assert headings[export.COLUMNS.index("country_of_performance")] == "Project Location"
    assert headings[export.COLUMNS.index("source")] == "Public Platform"
    assert headings[export.COLUMNS.index("notice_stage")] == "Phase of Tender"
    # A column the vocabulary does not name keeps its key rather than gaining an
    # invented label.
    assert headings[export.COLUMNS.index("confidence")] == "confidence"


def test_the_csv_carries_a_bom_so_excel_reads_accents() -> None:
    first = next(iter(export.to_csv([])))

    assert first == b"\xef\xbb\xbf"


def test_the_filename_is_dated() -> None:
    name = export.filename("watchdog-register", "csv", now=datetime(2026, 9, 13, tzinfo=UTC))

    assert name == "watchdog-register-2026-09-13.csv"


def _sheet(rows: list[RegisterRow]) -> Any:
    """The exported workbook, reopened, so the tests read what Excel would."""
    return load_workbook(io.BytesIO(export.to_xlsx(rows)))["Register"]


def _cell(sheet: Any, column: str) -> Any:
    return sheet.cell(row=2, column=export.COLUMNS.index(column) + 1)
